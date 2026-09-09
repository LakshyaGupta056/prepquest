"""
Helpers for pulling plain text out of uploaded files that Gemini cannot
read natively as inline bytes.

Gemini's file/image understanding handles PDFs, PNGs, and JPEGs directly,
so those are passed straight through as inline parts (see gemini_service.py).
.docx files are not natively understood, so we extract their text here
with python-docx before it ever reaches Gemini.
"""

import io

from docx import Document


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract all paragraph (and table cell) text from a .docx file's raw bytes.

    Raises:
        ValueError: if the bytes cannot be parsed as a .docx file.
    """
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:  # python-docx raises a range of exceptions on bad input
        raise ValueError(f"Could not read .docx file: {exc}") from exc

    chunks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))

    extracted = "\n".join(chunks).strip()

    if not extracted:
        raise ValueError("The uploaded .docx file appears to contain no readable text.")

    return extracted
